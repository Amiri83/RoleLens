"""Web flows: guards, accounts, collection job, roles, graph, editor, preview, save, export.

Placeholder data only (account 123456789012, role-a, bucket-a, fake credentials)."""

import html
import io
import json
import zipfile

import boto3
import pytest
from moto import mock_aws
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from rolelens.aws import AwsGateway
from rolelens.db import closing
from rolelens.filters import RoleFilters
from rolelens.web import create_app, host_allowed
from rolelens.xlsx import FINDING_HEADERS, HEADERS, roles_workbook

from .conftest import INLINE_A, seed_iam
from .test_graph_export import view_payload

FAKE_KEY_ID = "AKIAEXAMPLEEXAMPLE00"
FAKE_SECRET = "example/secret/value/for/tests/only/0000"


@pytest.fixture
def app(home):
    return create_app(home, testing=True, gateway_factory=lambda _acct: _gateway())


def _gateway() -> AwsGateway:
    return AwsGateway(boto3.session.Session(region_name="us-east-1"))


@pytest.fixture
def client(app):
    c = app.test_client()
    c.get("/")
    return c


def _csrf(client) -> str:
    with client.session_transaction() as s:
        return s["csrf"]


def _post(client, url, data=None, **kw):
    data = dict(data or {})
    data["csrf_token"] = _csrf(client)
    return client.post(url, data=data, **kw)


def _db(app):
    return app.extensions["rolelens"]["paths"].db_path


def _add_account(client, **fields) -> int:
    data = {"display_name": "example", "region": "us-east-1", "auth_mode": "env", **fields}
    resp = _post(client, "/accounts/new", data)
    assert resp.status_code == 302, resp.data
    return max(a.id for a in client.application.extensions["rolelens"]["accounts"].list())


def _collect(client) -> None:
    """Run the collection job in the request (no background=1: the plain form path)."""
    _add_account(client)
    resp = _post(client, "/collect")
    assert resp.status_code == 302
    page = client.get("/").get_data(as_text=True)
    assert "Collected example (123456789012): 2 roles" in page


def _draft_id(app) -> int:
    with closing(_db(app)) as conn:
        return conn.execute("SELECT MAX(id) FROM drafts").fetchone()[0]


# -- guards ---------------------------------------------------------------------------


def test_host_allowed():
    assert host_allowed("127.0.0.1:8078", 8078)
    assert host_allowed("localhost:8078", 8078)
    assert not host_allowed("example.com:8078", 8078)
    assert not host_allowed("127.0.0.1:9999", 8078)
    assert host_allowed("localhost", None)


def test_untrusted_host_header_is_rejected(home):
    app = create_app(home, testing=True, port=8078)
    c = app.test_client()
    assert c.get("/", headers={"Host": "attacker.example:8078"}).status_code == 400
    assert c.get("/", headers={"Host": "127.0.0.1:8078"}).status_code == 200


def test_post_without_csrf_token_is_rejected(client):
    assert client.post("/collect", data={}).status_code == 400
    assert client.post("/accounts/new", data={"auth_mode": "env"}).status_code == 400


# -- accounts -------------------------------------------------------------------------


def test_key_secret_is_encrypted_at_rest_and_never_rendered(app, client):
    _add_account(client, auth_mode="keys", access_key_id=FAKE_KEY_ID, secret_access_key=FAKE_SECRET)
    db = _db(app)
    for path in db.parent.glob(db.name + "*"):  # the database and its WAL file
        assert FAKE_SECRET.encode() not in path.read_bytes()
    with closing(_db(app)) as conn:
        row = conn.execute("SELECT secret_enc FROM accounts").fetchone()
    assert row["secret_enc"] and FAKE_SECRET not in row["secret_enc"]
    for url in ("/settings", "/accounts/1/edit"):
        page = client.get(url).get_data(as_text=True)
        assert FAKE_SECRET not in page and FAKE_KEY_ID not in page
    acct = app.extensions["rolelens"]["accounts"].get(1, with_secret=True)
    assert acct.secret_access_key == FAKE_SECRET


@mock_aws
def test_account_connection_test(client):
    _add_account(client)
    resp = _post(client, "/accounts/1/test", follow_redirects=True)
    assert "Connected to account 123456789012" in resp.get_data(as_text=True)


# -- collection, roles and graph --------------------------------------------------------


def test_roles_page_without_account_or_data(client):
    page = client.get("/").get_data(as_text=True)
    assert "No AWS account yet" in page
    assert client.get("/roles/role-a").status_code == 404


@mock_aws
def test_collect_then_roles_page_and_search(client):
    seed_iam()
    _collect(client)
    page = client.get("/").get_data(as_text=True)
    assert "role-a" in page and "role-b" in page
    assert "lambda.amazonaws.com" in page and "account 123456789012" in page
    searched = client.get("/?q=lambda").get_data(as_text=True)
    assert 'href="/roles/role-a"' in searched and 'href="/roles/role-b"' not in searched


@mock_aws
def test_collect_as_background_job(app, client):
    seed_iam()
    _add_account(client)
    resp = _post(client, "/collect", {"background": "1"})
    body = resp.get_json()
    assert body["ok"]
    assert app.extensions["rolelens"]["jobs"].wait(body["job"]["id"], 30)
    job = client.get(body["status_url"]).get_json()["job"]
    assert job["state"] == "done"
    assert any("2 roles" in m["text"] for m in job["messages"])


@mock_aws
def test_role_page_and_graph_json(client):
    seed_iam()
    _collect(client)
    page = client.get("/roles/role-a").get_data(as_text=True)
    assert "vendor/cytoscape.min.js" in page and "vendor/dagre.min.js" in page
    assert '<option value="grid">Grid</option>' in page  # Grid is the default (first) layout
    graph = client.get("/roles/role-a/graph.json").get_json()
    ids = {e["data"]["id"] for e in graph["elements"]}
    assert {"role", "policy:inline:policy-inline-a", "action:Deny:s3:deleteobject"} <= ids
    assert client.get("/roles/no-such-role").status_code == 404


# -- drafts ---------------------------------------------------------------------------


@mock_aws
def test_clone_edit_preview_save_and_export(app, client):
    seed_iam()
    _collect(client)
    resp = _post(client, "/roles/role-a/clone")
    assert resp.status_code == 302
    draft_id = _draft_id(app)
    editor = client.get(f"/drafts/{draft_id}").get_data(as_text=True)
    assert "never written to AWS" in editor
    assert "vendor/cytoscape.min.js" in editor and "jsonedit.js" in editor

    with closing(_db(app)) as conn:
        model = json.loads(conn.execute("SELECT model FROM drafts").fetchone()["model"])
    inline = next(p for p in model["policies"] if p["kind"] == "inline")

    # preview: an unknown action is reported and the graph / diff follow the edit
    doc = json.loads(json.dumps(INLINE_A))
    doc["Statement"][0]["Action"] = "s3:GetObjekt"
    inline["document"] = json.dumps(doc, indent=2)
    result = _post(client, f"/drafts/{draft_id}/preview", {"model": json.dumps(model)}).get_json()
    assert not result["ok"]
    assert any("unknown action 's3:GetObjekt'" in i["message"] for i in result["issues"])
    ids = {e["data"]["id"] for e in result["graph"]["elements"]}
    assert "action:Allow:s3:getobjekt" in ids
    assert result["diff"]["changed"] == 1

    # a JSON syntax error: no graph, the error names the document and line
    inline["document"] = '{\n  "Version": "2012-10-17",\n  "Statement": [\n}'
    result = _post(client, f"/drafts/{draft_id}/preview", {"model": json.dumps(model)}).get_json()
    assert result["graph"] is None
    assert result["issues"][0]["where"] == "policy:policy-inline-a"
    saved = _post(client, f"/drafts/{draft_id}", {"model": json.dumps(model), "ajax": "1"})
    assert saved.status_code == 400 and not saved.get_json()["saved"]

    # a valid edit is saved
    doc["Statement"][0]["Action"] = ["s3:GetObject", "s3:ListBucket"]
    inline["document"] = doc
    saved = _post(client, f"/drafts/{draft_id}", {"model": json.dumps(model), "ajax": "1"})
    assert saved.get_json()["saved"] and saved.get_json()["ok"]
    assert "1 section changed" in client.get("/drafts").get_data(as_text=True)

    export = html.unescape(client.get(f"/drafts/{draft_id}/export").get_data(as_text=True))
    assert 'resource "aws_iam_role" "role_a"' in export
    assert "s3:ListBucket" in export
    tf = client.get(f"/drafts/{draft_id}/export/main.tf")
    assert tf.headers["Content-Disposition"] == 'attachment; filename="role_a.tf"'
    assert b'resource "aws_iam_role_policy" "role_a_policy_inline_a"' in tf.data
    zipped = client.get(f"/drafts/{draft_id}/export/policies.zip")
    with zipfile.ZipFile(io.BytesIO(zipped.data)) as zf:
        assert "role_a/trust-policy.json" in zf.namelist()
    one = client.get(f"/drafts/{draft_id}/export/file?path=trust-policy.json")
    assert json.loads(one.data)["Statement"][0]["Principal"] == {"Service": "lambda.amazonaws.com"}
    assert client.get(f"/drafts/{draft_id}/export/file?path=../secret").status_code == 404

    # reset and delete
    _post(client, f"/drafts/{draft_id}/reset")
    assert "No changes yet" in client.get(f"/drafts/{draft_id}").get_data(as_text=True)
    _post(client, f"/drafts/{draft_id}/delete")
    assert client.get(f"/drafts/{draft_id}").status_code == 404


@mock_aws
def test_save_without_javascript_rerenders_the_raw_model_on_error(app, client):
    seed_iam()
    _collect(client)
    _post(client, "/roles/role-b/clone")
    draft_id = _draft_id(app)
    resp = _post(client, f"/drafts/{draft_id}", {"model": "{broken"})
    assert resp.status_code == 400
    assert "{broken" in resp.get_data(as_text=True)


@mock_aws
def test_drafts_never_call_aws_writes(app, client, monkeypatch):
    """Editing / exporting only touches SQLite: any AWS client creation would fail."""
    seed_iam()
    _collect(client)
    _post(client, "/roles/role-a/clone")
    draft_id = _draft_id(app)

    def no_aws(*_a, **_k):
        raise AssertionError("drafts must not call AWS")

    monkeypatch.setattr(AwsGateway, "client", no_aws)
    with closing(_db(app)) as conn:
        model = conn.execute("SELECT model FROM drafts").fetchone()["model"]
    assert _post(client, f"/drafts/{draft_id}", {"model": model, "ajax": "1"}).get_json()["saved"]
    assert client.get(f"/drafts/{draft_id}/export/main.tf").status_code == 200


# -- roles filters and Excel export -----------------------------------------------------


def _role_links(page: str) -> list[str]:
    return [name for name in ("role-a", "role-b") if f'href="/roles/{name}"' in page]


@mock_aws
def test_roles_column_filters(client):
    seed_iam()
    _collect(client)
    cases = {
        "": ["role-a", "role-b"],
        "trust=service": ["role-a"],
        "trust=account": ["role-b"],
        "trust=saml": [],
        "inline=yes": ["role-a"],
        "pmin=1": ["role-a"],
        "pmax=0": ["role-b"],
        "used=never": ["role-a", "role-b"],  # moto roles were never used
        "used=90": [],
        "slr=yes": [],
        "boundary=no": ["role-a", "role-b"],
        "q=/service/&trust=service": ["role-a"],
        "trust=bogus&pmin=x": ["role-a", "role-b"],  # unknown values are ignored
    }
    for query, expected in cases.items():
        page = client.get(f"/?{query}").get_data(as_text=True)
        assert _role_links(page) == expected, query
    page = client.get("/?trust=service").get_data(as_text=True)
    assert "1 of 2 roles match the filters" in page
    assert 'href="/roles.xlsx?trust=service"' in page  # the export follows the view
    assert '<option value="service" selected>' in page
    assert 'form="role-filters"' in page and 'name="pmin"' in page


@mock_aws
def test_roles_excel_export_of_the_filtered_view(client):
    seed_iam()
    _collect(client)
    resp = client.get("/roles.xlsx?trust=service&inline=yes")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resp.headers["Content-Disposition"].startswith('attachment; filename="rolelens-roles-')
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb["Roles"]
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == HEADERS
    assert [r[0] for r in rows[1:]] == ["role-a"]  # only what the filters keep
    row = dict(zip(HEADERS, rows[1], strict=True))
    assert row["Path"] == "/service/" and row["Trust types"] == "service"
    assert row["Trusted principals"] == "lambda.amazonaws.com"
    assert (row["Managed policies"], row["Inline policies"], row["Policies"]) == (1, 1, 2)
    assert row["Service-linked"] == "no" and row["Last used (UTC)"] is None
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == f"A1:{get_column_letter(len(HEADERS))}{len(rows)}"
    assert all(c.font.bold for c in ws[1])
    assert row["Highest severity"] == "low"
    assert row["Findings"] == "No permissions boundary (low)"
    findings = list(wb["Findings"].iter_rows(values_only=True))
    assert list(findings[0]) == FINDING_HEADERS
    assert [(f[0], f[1], f[3]) for f in findings[1:]] == [("role-a", "low", "no-boundary")]
    info = {r[0]: r[1] for r in wb["Export"].iter_rows(values_only=True)}
    assert info["Account"] == "example (123456789012)"
    assert info["Roles"] == "1 of 2"
    assert info["Findings"] == 1
    assert info["Trust type"] == "service" and info["Inline policies"] == "yes"

    everything = load_workbook(io.BytesIO(client.get("/roles.xlsx").data))
    assert [r[0] for r in everything["Roles"].iter_rows(min_row=2, values_only=True)] == [
        "role-a",
        "role-b",
    ]
    assert everything["Export"]["B6"].value == "none"  # Filters: none


def test_roles_excel_export_without_data(client):
    assert client.get("/roles.xlsx").status_code == 404


def test_excel_cells_are_never_formulas():
    row = {
        "name": "role-a",
        "path": "/",
        "arn": "arn:aws:iam::123456789012:role/role-a",
        "description": '=HYPERLINK("http://example.invalid","x")\x07',
        "last_used_at": "2026-05-01T10:00:00+00:00",
        "last_used_region": "us-east-1",
        "last_used_days": 31,
        "principals": ["lambda.amazonaws.com"],
        "trust_types": ["service"],
        "service_linked": False,
        "boundary": "",
        "managed": 0,
        "inline": 1,
        "policies": 1,
    }
    data = roles_workbook([row], RoleFilters(), account="123456789012", taken_at="", total=1)
    ws = load_workbook(io.BytesIO(data))["Roles"]
    cell = ws["D2"]
    assert cell.data_type == "s" and cell.value == '=HYPERLINK("http://example.invalid","x")'
    assert ws["E2"].value.isoformat() == "2026-05-01T10:00:00"  # naive UTC
    assert ws["G2"].value == 31


# -- graph exports and the job modal ------------------------------------------------------


@mock_aws
def test_graph_export_endpoint(client):
    seed_iam()
    _collect(client)
    page = client.get("/roles/role-a").get_data(as_text=True)
    assert 'data-export-url="/graph/export.__FMT__"' in page
    assert 'id="graph-icons"' in page and "data:image/svg+xml;base64," in page
    for fmt in ("svg", "png", "drawio"):
        assert f'data-export="{fmt}"' in page

    graph = client.get("/roles/role-a/graph.json").get_json()
    view = json.dumps(view_payload(graph, keep={"role", "policy", "service"}))
    svg = _post(client, "/graph/export.svg", {"view": view})
    assert svg.status_code == 200 and svg.mimetype == "image/svg+xml"
    assert svg.headers["Content-Disposition"] == 'attachment; filename="role_a-graph.svg"'
    assert b"s3:GetObject" not in svg.data  # actions were filtered out of the view
    drawio = _post(client, "/graph/export.drawio", {"view": view})
    assert drawio.status_code == 200 and drawio.data.startswith(b"<?xml")
    assert b"<mxfile" in drawio.data and b"mxgraph.aws4.s3" in drawio.data
    assert _post(client, "/graph/export.svg", {"view": "{}"}).status_code == 400
    assert _post(client, "/graph/export.pdf", {"view": view}).status_code == 404
    assert client.post("/graph/export.svg", data={"view": view}).status_code == 400  # no CSRF


@mock_aws
def test_job_modal_close_and_data_refresh_markup(app, client):
    seed_iam()
    _add_account(client)
    page = client.get("/").get_data(as_text=True)
    assert 'id="job-close">Close</button>' in page
    assert "job-hide" not in page and "Close and reload" not in page
    assert "refresh.js" in page and 'id="roles-data" data-refresh' in page
    assert 'id="header-ctx" data-refresh' in page
    body = _post(client, "/collect", {"background": "1"}).get_json()
    assert app.extensions["rolelens"]["jobs"].wait(body["job"]["id"], 30)
    job = client.get(body["status_url"]).get_json()["job"]
    assert job["changes_data"] is True
    assert client.get("/jobs/current").get_json()["job"]["changes_data"] is True


def test_logs_page(client):
    assert client.get("/logs").status_code == 200
    assert client.get("/settings").status_code == 200
