"""Web flows of users, groups and managed policies: pages, filters, Excel exports, graphs
with click-through links, recommendations, drafts and their JSON / Terraform export.

Placeholder data only (user-a, group-a, policy-a, role-a, 123456789012, bucket-a)."""

import io
import json
import zipfile
from urllib.parse import quote

import pytest
from moto import mock_aws
from openpyxl import load_workbook

from rolelens import export
from rolelens.xlsx import POLICY_HEADERS, USER_HEADERS

from .conftest import ADMIN_DOC, POLICY_A, seed_iam, seed_identities
from .test_web import _add_account, _draft_id, _post, app, client  # noqa: F401 - fixtures

POLICY_ARN = "arn:aws:iam::123456789012:policy/policy-a"
UNUSED_ARN = "arn:aws:iam::123456789012:policy/policy-unused"


@pytest.fixture
def collected(client):  # noqa: F811 - pytest fixture
    with mock_aws():
        arns = seed_iam()
        seed = seed_identities(arns["policy_arn"])
        _add_account(client)
        assert _post(client, "/collect").status_code == 302
        page = client.get("/").get_data(as_text=True)
        assert "Collected example (123456789012): 2 roles, 4 users, 2 groups" in page
        yield seed


def _sheet_rows(data: bytes, sheet: str) -> list[list]:
    wb = load_workbook(io.BytesIO(data))
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]


def test_pages_without_data(client):  # noqa: F811
    for url in ("/users", "/groups", "/policies"):
        resp = client.get(url)
        assert resp.status_code == 200, url
        assert "No AWS account yet" in resp.get_data(as_text=True)
    assert client.get("/users.xlsx").status_code == 404
    assert client.get("/policies.xlsx").status_code == 404
    assert client.get("/users/user-a").status_code == 404


# -- users ---------------------------------------------------------------------------------


def test_users_page_and_filters(client, collected):  # noqa: F811
    page = client.get("/users").get_data(as_text=True)
    for name in ("user-a", "user-b", "user-c", "user-admin"):
        assert f'href="/users/{name}"' in page
    assert "no MFA" in page
    assert collected["key_id"] not in page  # only the shortened key id is ever shown

    page = client.get("/users?mfa=no&console=yes").get_data(as_text=True)
    assert "1 of 4 users match the filters" in page
    assert 'href="/users/user-a"' in page and 'href="/users/user-b"' not in page
    page = client.get("/users?perms=group").get_data(as_text=True)
    assert "3 of 4 users match" in page and 'href="/users/user-a"' not in page
    page = client.get("/users?finding=admin").get_data(as_text=True)
    assert 'href="/users/user-admin"' in page and 'href="/users/user-b"' not in page
    # the export link carries the filters
    assert 'href="/users.xlsx?finding=admin"' in page


def test_users_excel_export_columns(client, collected):  # noqa: F811
    resp = client.get("/users.xlsx?inline=yes")
    assert resp.status_code == 200
    assert "rolelens-users-" in resp.headers["Content-Disposition"]
    rows = _sheet_rows(resp.data, "Users")
    assert rows[0] == USER_HEADERS
    assert len(rows) == 2  # header + user-a
    user = dict(zip(USER_HEADERS, rows[1], strict=True))
    assert user["User"] == "user-a"
    assert user["Groups"] == "group-a"
    assert user["Console access"] == "yes" and user["MFA"] == "no"
    assert user["Active access keys"] == 1
    assert user["Direct inline policies"] == 1 and user["Policies via groups"] == 2
    assert user["Permissions from"] == "direct and via groups"
    assert "Console access without MFA (high)" in user["Findings"]
    findings = _sheet_rows(resp.data, "Findings")
    assert findings[0][0] == "User"
    assert {r[3] for r in findings[1:]} >= {"no-mfa", "direct-attachment"}
    info = dict((r[0], r[1]) for r in _sheet_rows(resp.data, "Export"))
    assert info["Users"] == "1 of 4" and info["Inline policies"] == "yes"
    # nothing secret in the workbook
    blob = json.dumps([_sheet_rows(resp.data, s) for s in ("Users", "Findings")], default=str)
    assert collected["secret"] not in blob and collected["key_id"] not in blob


def test_user_page_graph_and_click_through(client, collected):  # noqa: F811
    page = client.get("/users/user-a").get_data(as_text=True)
    assert "Console access without MFA" in page
    assert "…" + collected["key_id"][-4:] in page and collected["key_id"] not in page
    assert "policy-a" in page and "via group" in page

    graph = client.get("/users/user-a/graph.json").get_json()
    nodes = {e["data"]["id"]: e["data"] for e in graph["elements"] if e["group"] == "nodes"}
    edges = {
        (e["data"]["source"], e["data"]["target"])
        for e in graph["elements"]
        if e["group"] == "edges"
    }
    assert nodes["user"]["url"] == "/users/user-a"
    assert nodes["group:group-a"]["url"] == "/groups/group-a"
    assert nodes["policy:managed:policy-a"]["url"] == "/policies/view?arn=" + POLICY_ARN
    assert ("user", "group:group-a") in edges
    assert ("group:group-a", "policy:managed:policy-a") in edges
    assert ("user", "policy:inline:user/user-inline-a") in edges
    assert ("policy:managed:policy-a", "service:sqs") in edges
    assert graph["stats"]["groups"] == 1 and graph["stats"]["users"] == 1

    # role graphs link their managed policies too
    role = client.get("/roles/role-a/graph.json").get_json()
    role_nodes = {e["data"]["id"]: e["data"] for e in role["elements"] if e["group"] == "nodes"}
    assert role_nodes["role"]["url"] == "/roles/role-a"
    assert role_nodes["policy:managed:policy-a"]["url"].startswith("/policies/view?arn=")
    assert "url" not in role_nodes["policy:inline:policy-inline-a"]
    assert client.get("/users/no-such-user").status_code == 404


def test_groups_pages(client, collected):  # noqa: F811
    page = client.get("/groups").get_data(as_text=True)
    assert 'href="/groups/group-a"' in page and 'href="/groups/group-admin"' in page
    assert "admin" in page  # group-admin's finding badge
    page = client.get("/groups?q=user-admin").get_data(as_text=True)
    assert "1 of 2 groups match" in page and 'href="/groups/group-a"' not in page
    page = client.get("/groups/group-a").get_data(as_text=True)
    assert "Members (3)" in page and 'href="/users/user-b"' in page
    graph = client.get("/groups/group-admin/graph.json").get_json()
    ids = {e["data"]["id"] for e in graph["elements"] if e["group"] == "nodes"}
    assert {"group:group-admin", "policy:inline:group:group-admin/admin-inline"} <= ids


# -- policies ------------------------------------------------------------------------------


def test_policies_page_usage_and_unattached(client, collected):  # noqa: F811
    page = client.get("/policies").get_data(as_text=True)
    assert "policy-a" in page and "policy-unused" in page
    assert page.count('badge warn">unattached') == 1
    page = client.get("/policies?unattached=yes").get_data(as_text=True)
    assert "1 of 2 policies match" in page and ">policy-a<" not in page
    page = client.get("/policies?used_by=group").get_data(as_text=True)
    assert ">policy-a<" in page and "policy-unused" not in page

    detail = client.get("/policies/view?arn=" + quote(POLICY_ARN, safe="")).get_data(as_text=True)
    assert 'href="/roles/role-a"' in detail and 'href="/groups/group-a"' in detail
    assert "sqs:SendMessage" in detail  # JSON view
    assert "v1" in detail and "default" in detail  # versions
    unused = client.get("/policies/view?arn=" + quote(UNUSED_ARN, safe="")).get_data(as_text=True)
    assert "Attached to no role, user or group" in unused
    assert client.get("/policies/view?arn=arn:aws:iam::123456789012:policy/none").status_code == 404


def test_policies_excel_export_columns(client, collected):  # noqa: F811
    resp = client.get("/policies.xlsx")
    assert resp.status_code == 200
    rows = _sheet_rows(resp.data, "Policies")
    assert rows[0] == POLICY_HEADERS
    by_name = {r[0]: dict(zip(POLICY_HEADERS, r, strict=True)) for r in rows[1:]}
    assert by_name["policy-a"]["Roles"] == "role-a"
    assert by_name["policy-a"]["Groups"] == "group-a"
    assert by_name["policy-a"]["Unattached"] == "no"
    assert by_name["policy-a"]["Default version"] == "v1"
    assert by_name["policy-unused"]["Unattached"] == "yes"
    assert by_name["policy-unused"]["Type"] == "customer managed"
    assert _sheet_rows(resp.data, "Findings")[0][0] == "Policy"


# -- insights ------------------------------------------------------------------------------


def test_insights_and_recommendations_cover_users_and_groups(client, collected):  # noqa: F811
    page = client.get("/insights").get_data(as_text=True)
    assert "Users and groups" in page and "(4 users, 2 groups)" in page
    assert "Console access without MFA" in page
    assert 'href="/users?finding=no-mfa"' in page
    assert 'href="/users/user-admin#findings"' in page

    page = client.get("/insights/recommendations").get_data(as_text=True)
    assert "Duplicate users (1 group)" in page
    start = page.index('id="user-duplicates"')
    table = page[start : page.index("</table>", start)]
    for name in ("user-a", "user-b", "user-c"):
        assert f'href="/users/{name}"' in table


# -- drafts and export ---------------------------------------------------------------------


def test_clone_user_edit_and_export(app, client, collected):  # noqa: F811
    resp = _post(client, "/users/user-a/clone")
    assert resp.status_code == 302
    draft_id = _draft_id(app)
    page = client.get(f"/drafts/{draft_id}").get_data(as_text=True)
    assert 'data-entity="user"' in page and 'data-field="groups"' in page
    assert 'data-field="role_name"' not in page

    model = {
        "entity": "user",
        "user_name": "user-a",
        "path": "/",
        "permissions_boundary": "",
        "groups": ["group-a", "group-b"],
        "policies": [
            {"name": "user-inline-a", "kind": "inline", "document": json.dumps(POLICY_A)},
            {"name": "admin-inline", "kind": "inline", "document": ADMIN_DOC},
        ],
    }
    preview = _post(client, f"/drafts/{draft_id}/preview", {"model": json.dumps(model)}).get_json()
    assert preview["ok"], preview["issues"]
    titles = {s["title"]: s["status"] for s in preview["diff"]["sections"]}
    assert titles["User settings"] == "changed"  # group-b added
    assert titles["admin-inline (inline)"] == "added"
    graph_ids = {e["data"]["id"] for e in preview["graph"]["elements"]}
    assert "group:group-b" in graph_ids

    bad = dict(model, user_name="user a!", groups=["group a!"])
    issues = _post(client, f"/drafts/{draft_id}/preview", {"model": json.dumps(bad)}).get_json()
    assert {(i["where"], i["path"]) for i in issues["issues"] if i["level"] == "error"} >= {
        ("user", "user_name"),
        ("user", "groups[0]"),
    }

    assert _post(client, f"/drafts/{draft_id}", {"model": json.dumps(model)}).status_code == 302
    tf = client.get(f"/drafts/{draft_id}/export/main.tf").get_data(as_text=True)
    assert 'resource "aws_iam_user" "user_a"' in tf
    assert 'resource "aws_iam_user_policy" "user_a_admin_inline"' in tf
    assert "user   = aws_iam_user.user_a.name" in tf
    assert 'groups = ["group-a", "group-b"]' in tf
    assert "cloned from user user-a" in tf
    assert "access_key" not in tf and "login_profile" not in tf
    data = client.get(f"/drafts/{draft_id}/export/policies.zip").data
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = sorted(zf.namelist())
        user_json = json.loads(zf.read("user_a/user.json"))
    assert names == [
        "user_a/policies/inline-admin-inline.json",
        "user_a/policies/inline-user-inline-a.json",
        "user_a/user.json",
    ]
    assert user_json["Groups"] == ["group-a", "group-b"]
    assert collected["key_id"] not in json.dumps(user_json)
    drafts = client.get("/drafts").get_data(as_text=True)
    assert "user-a" in drafts and ">user<" in drafts


def test_clone_managed_policy_edit_and_export(app, client, collected):  # noqa: F811
    resp = _post(client, "/policies/clone", {"arn": POLICY_ARN})
    assert resp.status_code == 302
    draft_id = _draft_id(app)
    page = client.get(f"/drafts/{draft_id}").get_data(as_text=True)
    assert 'data-entity="policy"' in page and 'id="doc-add"' not in page

    model = {
        "entity": "policy",
        "name": "policy-a",
        "path": "/",
        "description": "Example policy",
        "document": json.dumps(dict(POLICY_A, Statement=[*POLICY_A["Statement"]])),
    }
    preview = _post(client, f"/drafts/{draft_id}/preview", {"model": json.dumps(model)}).get_json()
    assert preview["ok"], preview["issues"]
    assert {s["title"]: s["status"] for s in preview["diff"]["sections"]}[
        "Policy settings"
    ] == "changed"
    broken = dict(model, document='{"Version": "2012-10-17", "Statement": [')
    resp = _post(client, f"/drafts/{draft_id}/preview", {"model": json.dumps(broken)}).get_json()
    assert any(
        i["where"] == "policy:policy-a" and "invalid JSON" in i["message"] for i in resp["issues"]
    )

    assert _post(client, f"/drafts/{draft_id}", {"model": json.dumps(model)}).status_code == 302
    tf = client.get(f"/drafts/{draft_id}/export/main.tf").get_data(as_text=True)
    assert 'resource "aws_iam_policy" "policy_a"' in tf
    assert 'description = "Example policy"' in tf
    assert f"terraform import aws_iam_policy.policy_a {POLICY_ARN}" in tf
    assert "aws_iam_role" not in tf
    files = export.json_files(_draft_model(app, draft_id))
    assert sorted(files) == ["policies/managed-policy-a.json", "policy.json"]
    assert json.loads(files["policies/managed-policy-a.json"]) == POLICY_A
    # AWS managed policies and unknown ARNs cannot be cloned
    assert (
        _post(client, "/policies/clone", {"arn": "arn:aws:iam::aws:policy/none"}).status_code == 404
    )


def _draft_model(app, draft_id):  # noqa: F811
    from rolelens.db import closing

    with closing(app.extensions["rolelens"]["paths"].db_path) as conn:
        row = conn.execute("SELECT model, entity FROM drafts WHERE id=?", (draft_id,)).fetchone()
    assert row["entity"] == "policy"
    return json.loads(row["model"])


def test_role_drafts_are_unchanged(app, client, collected):  # noqa: F811
    assert _post(client, "/roles/role-a/clone").status_code == 302
    draft_id = _draft_id(app)
    page = client.get(f"/drafts/{draft_id}").get_data(as_text=True)
    assert 'data-entity="role"' in page and 'data-field="role_name"' in page
    tf = client.get(f"/drafts/{draft_id}/export/main.tf").get_data(as_text=True)
    assert 'resource "aws_iam_role" "role_a"' in tf and "cloned from role role-a" in tf
