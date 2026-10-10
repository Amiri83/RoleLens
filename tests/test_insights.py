"""Insights pages end to end: analysis at collection, dashboard, roles filter by finding,
role page sections, recommendations, compare / merge exports, the usage job, settings.

moto for IAM, botocore Stubber for the usage sources. Placeholder data only (account
123456789012, role-a ... role-d, bucket-a, documentation IP 192.0.2.10)."""

import io
import json
import zipfile
from datetime import UTC, datetime, timedelta

import boto3
import pytest
from moto import mock_aws
from openpyxl import load_workbook

from rolelens.db import closing
from rolelens.insights import ANALYSIS_VERSION
from rolelens.usage import ASSUME_EVENTS
from rolelens.web import create_app

from .conftest import ACCOUNT, LAMBDA_TRUST, seed_iam
from .test_usage import IP, StubGateway, _job, _record, _service
from .test_web import _add_account, _db, _gateway, _post

READ_BUCKET = {
    "Version": "2012-10-17",
    "Statement": [
        {"Effect": "Allow", "Action": "s3:GetObject", "Resource": "arn:aws:s3:::bucket-a/*"}
    ],
}
ADMIN = {
    "Version": "2012-10-17",
    "Statement": [{"Effect": "Allow", "Action": "*", "Resource": "*"}],
}
ROLES = ["role-a", "role-admin", "role-b", "role-c", "role-d"]


@pytest.fixture
def app(home):
    return create_app(home, testing=True, gateway_factory=lambda _acct: _gateway())


@pytest.fixture
def client(app):
    c = app.test_client()
    c.get("/")
    return c


def _seed():
    """seed_iam() (role-a, role-b) plus duplicates role-c / role-d and role-admin."""
    seed_iam()
    iam = boto3.client("iam", region_name="us-east-1")
    for name in ("role-c", "role-d"):
        iam.create_role(RoleName=name, AssumeRolePolicyDocument=json.dumps(LAMBDA_TRUST))
        iam.put_role_policy(
            RoleName=name, PolicyName="read-bucket-a", PolicyDocument=json.dumps(READ_BUCKET)
        )
    iam.create_role(RoleName="role-admin", AssumeRolePolicyDocument=json.dumps(LAMBDA_TRUST))
    iam.put_role_policy(RoleName="role-admin", PolicyName="admin", PolicyDocument=json.dumps(ADMIN))


def _collect(client):
    _seed()
    _add_account(client)
    assert _post(client, "/collect").status_code == 302
    page = client.get("/").get_data(as_text=True)
    assert "Collected example (123456789012): 5 roles" in page


def _links(page):
    return [n for n in ROLES if f'href="/roles/{n}"' in page]


@mock_aws
def test_collection_stores_effective_permissions_and_findings(app, client):
    _collect(client)
    with closing(_db(app)) as conn:
        assert conn.execute("SELECT version FROM analysis").fetchone()[0] == ANALYSIS_VERSION
        actions = {
            r["action"]
            for r in conn.execute("SELECT action FROM role_permissions WHERE role_name='role-a'")
        }
        rules = {
            (r["role_name"], r["rule"], r["severity"])
            for r in conn.execute("SELECT * FROM role_findings")
        }
    assert actions == {"s3:GetObject", "sqs:SendMessage"}  # the Deny is not a grant
    assert ("role-admin", "admin", "critical") in rules
    assert ("role-admin", "passrole", "high") in rules
    assert ("role-c", "inline-only", "low") in rules
    assert ("role-a", "inline-only", "low") not in rules  # role-a has a managed policy too


@mock_aws
def test_insights_dashboard_and_roles_filter_by_finding(client):
    page = client.get("/insights").get_data(as_text=True)
    assert "No AWS account yet" in page  # no data yet
    _collect(client)
    page = client.get("/insights").get_data(as_text=True)
    assert 'href="/insights"' in page and 'aria-current="page"' in page
    assert 'id="rule-counts"' in page and "Admin-equivalent permissions" in page
    assert 'href="/?finding=admin"' in page
    assert 'id="risky-roles"' in page and 'href="/roles/role-admin#findings"' in page
    assert "Needs service last accessed data" in page
    assert 'data-job="usage"' in page

    cases = {
        "finding=admin": ["role-admin"],
        "finding=inline-only": ["role-admin", "role-c", "role-d"],
        "finding=no-boundary": ROLES,
        "finding=any": ROLES,
        "finding=none": [],
        "finding=bogus": ROLES,  # unknown values are ignored
        "finding=inline-only&q=role-c": ["role-c"],
    }
    for query, expected in cases.items():
        assert _links(client.get(f"/?{query}").get_data(as_text=True)) == expected, query
    page = client.get("/?finding=admin").get_data(as_text=True)
    assert 'class="badge sev-critical"' in page and '<option value="admin" selected>' in page
    assert 'href="/roles.xlsx?finding=admin"' in page

    wb = load_workbook(io.BytesIO(client.get("/roles.xlsx?finding=admin").data))
    rows = list(wb["Roles"].iter_rows(values_only=True))
    assert [r[0] for r in rows[1:]] == ["role-admin"]
    findings = list(wb["Findings"].iter_rows(min_row=2, values_only=True))
    assert {f[0] for f in findings} == {"role-admin"}
    assert {f[3] for f in findings} >= {"admin", "passrole", "inline-only", "no-boundary"}
    info = {r[0]: r[1] for r in wb["Export"].iter_rows(values_only=True)}
    assert info["Finding"] == "Admin-equivalent permissions"


@mock_aws
def test_role_page_shows_findings_effective_permissions_and_usage_hint(client):
    _collect(client)
    page = client.get("/roles/role-admin").get_data(as_text=True)
    assert 'id="role-insights" data-refresh' in page
    assert "Admin-equivalent permissions" in page and "Suggested fix" in page
    assert 'class="badge sev-critical"' in page
    assert 'id="effective-permissions"' in page and "s3:GetObject" in page
    assert "No usage data for this snapshot" in page
    page = client.get("/roles/role-b").get_data(as_text=True)
    assert "No permissions boundary" in page and "0 actions" in page


@mock_aws
def test_older_snapshots_are_analysed_on_first_view(app, client):
    _collect(client)
    with closing(_db(app)) as conn:
        for table in ("analysis", "role_permissions", "role_findings"):
            conn.execute(f"DELETE FROM {table}")  # noqa: S608 - fixed table names
    assert _links(client.get("/?finding=admin").get_data(as_text=True)) == ["role-admin"]
    with closing(_db(app)) as conn:
        assert conn.execute("SELECT COUNT(*) FROM role_findings").fetchone()[0] > 0


@mock_aws
def test_recommendations_compare_and_merge(app, client):
    _collect(client)
    page = client.get("/insights/recommendations").get_data(as_text=True)
    assert 'id="duplicates"' in page
    assert 'href="/insights/compare?a=role-c&amp;b=role-d"' in page
    # role-a (s3:GetObject + sqs:SendMessage) vs role-c (s3:GetObject): Jaccard 0.5
    assert "/insights/compare?a=role-a&amp;b=role-c" not in page
    page = client.get("/insights/recommendations?threshold=0.5").get_data(as_text=True)
    assert 'href="/insights/compare?a=role-a&amp;b=role-c"' in page
    assert "A ⊃ B" in page and 'value="0.5"' in page

    page = client.get("/insights/compare?a=role-a&b=role-c").get_data(as_text=True)
    assert 'id="side-by-side"' in page and "sqs:SendMessage" in page
    assert "role-a-merged" in page and "merged-inline" in page
    tf = client.get("/insights/merge.tf?a=role-a&b=role-c")
    assert tf.status_code == 200
    assert tf.headers["Content-Disposition"] == 'attachment; filename="role_a_merged.tf"'
    body = tf.get_data(as_text=True)
    assert 'resource "aws_iam_role" "role_a_merged"' in body
    assert 'aws_iam_role_policy" "role_a_merged_merged_inline"' in body
    assert "suggested merge of roles role-a and role-c" in body
    zf = zipfile.ZipFile(io.BytesIO(client.get("/insights/merge.zip?a=role-a&b=role-c").data))
    inline = json.loads(zf.read("role_a_merged/policies/inline-merged-inline.json"))
    # role-a's and role-c's identical s3:GetObject statements (Sid aside) are kept once
    assert [(st["Effect"], st["Action"]) for st in inline["Statement"]] == [
        ("Allow", "s3:GetObject"),
        ("Deny", "s3:DeleteObject"),
    ]
    assert "role_a_merged/role.json" in zf.namelist()

    assert client.get("/insights/compare?a=role-a&b=role-a").status_code == 400
    assert client.get("/insights/compare?a=role-a&b=no-such-role").status_code == 404
    assert client.post("/insights/merge/draft?a=role-a&b=role-c").status_code == 400  # no CSRF
    resp = _post(client, "/insights/merge/draft?a=role-a&b=role-c")
    assert resp.status_code == 302 and "/drafts/" in resp.headers["Location"]
    with closing(_db(app)) as conn:
        draft = conn.execute("SELECT * FROM drafts").fetchone()
    assert draft["source_role"] == "role-a"
    assert json.loads(draft["model"])["role_name"] == "role-a-merged"
    assert json.loads(draft["original"])["role_name"] == "role-a"


def _usage_gateway():
    """Stubbed usage sources for the five collected roles (role name order)."""
    gw = StubGateway()
    iam = gw.stub("iam")
    for i, _name in enumerate(ROLES):
        job = f"00000000-0000-0000-0000-00000000000{i}"
        iam.add_response("generate_service_last_accessed_details", {"JobId": job}, None)
    now = datetime.now(UTC)
    for name in ROLES:
        services = []
        if name == "role-a":
            services = [_service("s3", now - timedelta(days=1)), _service("sqs")]
        if name == "role-admin":
            services = [_service("ec2", now - timedelta(days=120))]
        iam.add_response("get_service_last_accessed_details", _job("COMPLETED", services), None)
    ct = gw.stub("cloudtrail")
    role_a = f"arn:aws:iam::{ACCOUNT}:role/service/role-a"
    for name in ASSUME_EVENTS:
        events = []
        if name == "AssumeRole":
            rec = _record(role_a, when=now - timedelta(hours=2))
            events = [{"EventId": "e1", "CloudTrailEvent": json.dumps(rec)}]
        ct.add_response("lookup_events", {"Events": events}, None)
    gw.stub("accessanalyzer").add_client_error("list_analyzers", "AccessDeniedException", "x")
    return gw


@mock_aws
def test_usage_job_fills_the_role_page_and_dashboard(app, client):
    _collect(client)
    app.extensions["rolelens"]["gateway_factory"] = lambda _acct: _usage_gateway()
    body = _post(client, "/insights/usage", {"background": "1"}).get_json()
    assert body["ok"] and body["job"]["kind"] == "usage"
    assert app.extensions["rolelens"]["jobs"].wait(body["job"]["id"], 30)
    job = client.get(body["status_url"]).get_json()["job"]
    assert job["state"] == "done" and job["changes_data"] is True
    texts = [m["text"] for m in job["messages"]]
    assert texts[0].startswith("Usage collected: 3 service entries, 1 AssumeRole events")
    assert any("IAM Access Analyzer" in t and "AccessDeniedException" in t for t in texts)

    page = client.get("/roles/role-a").get_data(as_text=True)
    assert 'id="service-access"' in page and "1 unused service" in page
    assert 'id="assume-events"' in page and IP in page and "role-ci" in page
    page = client.get("/insights").get_data(as_text=True)
    assert 'id="top-unused"' in page and 'id="unused-roles"' in page
    assert "unavailable (AccessDeniedException)" in page
    assert 'id="analyzer-findings"' not in page


def test_usage_job_needs_collected_data(client):
    resp = _post(client, "/insights/usage", {"background": "1"})
    assert resp.status_code == 400
    assert "Collect the IAM roles" in resp.get_json()["error"]


def test_overlap_threshold_setting(app, client):
    assert (
        _post(client, "/settings", {"log_dir": "", "overlap_threshold": "0.9"}).status_code == 302
    )
    store = app.extensions["rolelens"]["store"]
    assert store.load().overlap_threshold == 0.9
    page = client.get("/settings").get_data(as_text=True)
    assert 'name="overlap_threshold"' in page and 'value="0.9"' in page
    _post(client, "/settings", {"log_dir": "", "overlap_threshold": "2"})
    assert store.load().overlap_threshold == 0.9  # rejected, unchanged
    assert "must be a number from 0.5 to 1" in client.get("/settings").get_data(as_text=True)
