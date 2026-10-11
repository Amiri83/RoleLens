"""Excel exports of the Roles, Users and Policies pages (openpyxl): the rows of the
current filtered view.

The first sheet ("Roles", "Users" or "Policies") has a bold, frozen header row and an
autofilter over every column; sheet "Findings" lists the risk findings of those rows
(severity, why, suggested fix); sheet "Export" records the account, the snapshot time and
the filters that were applied. Every text cell is written as a string, so a value such as
``=cmd`` taken from a description can never become a formula.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Mapping, Sequence
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .filters import PERMISSION_SOURCES, PolicyFilters, RoleFilters, UserFilters
from .findings import RULES
from .policy import TRUST_TYPES

DATE_FORMAT = "yyyy-mm-dd hh:mm"
HEADER_FILL = PatternFill("solid", fgColor="E8EDF3")


def _utc_naive(value: str) -> datetime | None:
    """Excel has no time zones: timestamps are written as naive UTC."""
    if not value:
        return None
    try:
        ts = datetime.fromisoformat(value)
    except ValueError:
        return None
    if ts.tzinfo is not None:
        ts = ts.astimezone(UTC).replace(tzinfo=None)
    return ts


def _yes_no(flag: Any) -> str:
    return "yes" if flag else "no"


def _known(flag: bool | None) -> str:
    return "unknown" if flag is None else _yes_no(flag)


def _findings(by_rule: Mapping[str, str]) -> str:
    """``"Title (severity), ..."`` of a row's findings (rule -> worst severity)."""
    return ", ".join(f"{RULES[k].title if k in RULES else k} ({v})" for k, v in by_rule.items())


Column = tuple[str, int, Callable[[Mapping[str, Any]], Any]]
COLUMNS: list[Column] = [
    ("Role", 32, lambda r: r["name"]),
    ("Path", 18, lambda r: r["path"]),
    ("ARN", 48, lambda r: r["arn"]),
    ("Description", 36, lambda r: r["description"]),
    ("Last used (UTC)", 18, lambda r: _utc_naive(r["last_used_at"])),
    ("Last used region", 14, lambda r: r["last_used_region"]),
    ("Days since last use", 12, lambda r: r["last_used_days"]),
    ("Trust types", 22, lambda r: ", ".join(TRUST_TYPES[t] for t in r["trust_types"])),
    ("Trusted principals", 40, lambda r: ", ".join(r["principals"])),
    ("Service-linked", 10, lambda r: _yes_no(r["service_linked"])),
    ("Permissions boundary", 40, lambda r: r["boundary"]),
    ("Managed policies", 10, lambda r: r["managed"]),
    ("Inline policies", 10, lambda r: r["inline"]),
    ("Policies", 10, lambda r: r["policies"]),
    ("Highest severity", 12, lambda r: r.get("severity", "")),
    ("Findings", 48, lambda r: _findings(r.get("findings") or {})),
]
HEADERS = [c[0] for c in COLUMNS]
USER_COLUMNS: list[Column] = [
    ("User", 28, lambda r: r["name"]),
    ("Path", 14, lambda r: r["path"]),
    ("ARN", 44, lambda r: r["arn"]),
    ("Created (UTC)", 18, lambda r: _utc_naive(r["create_date"])),
    ("Groups", 30, lambda r: ", ".join(r["groups"])),
    ("Console access", 10, lambda r: _known(r["console"])),
    ("Console last used (UTC)", 18, lambda r: _utc_naive(r["password_last_used"])),
    ("MFA", 10, lambda r: _known(r["mfa"])),
    ("Active access keys", 10, lambda r: r["active_keys"]),
    ("Oldest active key (days)", 12, lambda r: r["key_age_days"]),
    ("Key unused (days)", 12, lambda r: r["key_unused_days"]),
    ("Last activity (UTC)", 18, lambda r: _utc_naive(r["last_activity"])),
    ("Days without activity", 12, lambda r: r["idle_days"]),
    ("Permissions boundary", 36, lambda r: r["boundary"]),
    ("Direct managed policies", 10, lambda r: r["managed"]),
    ("Direct inline policies", 10, lambda r: r["inline"]),
    ("Policies via groups", 10, lambda r: r["group_policies"]),
    ("Permissions from", 18, lambda r: PERMISSION_SOURCES.get(r["permissions"], r["permissions"])),
    ("Highest severity", 12, lambda r: r.get("severity", "")),
    ("Findings", 48, lambda r: _findings(r.get("findings") or {})),
]
USER_HEADERS = [c[0] for c in USER_COLUMNS]
POLICY_COLUMNS: list[Column] = [
    ("Policy", 32, lambda r: r["name"]),
    ("Type", 16, lambda r: "AWS managed" if r["aws_managed"] else "customer managed"),
    ("Path", 14, lambda r: r["path"]),
    ("ARN", 52, lambda r: r["arn"]),
    ("Description", 36, lambda r: r["description"]),
    ("Default version", 10, lambda r: r["default_version"]),
    ("Versions", 10, lambda r: r["version_count"]),
    ("Created (UTC)", 18, lambda r: _utc_naive(r["create_date"])),
    ("Updated (UTC)", 18, lambda r: _utc_naive(r["update_date"])),
    ("Attachment count", 10, lambda r: r["attachment_count"]),
    ("Unattached", 10, lambda r: _yes_no(r["unattached"])),
    ("Roles", 36, lambda r: ", ".join(r["roles"])),
    ("Users", 36, lambda r: ", ".join(r["users"])),
    ("Groups", 36, lambda r: ", ".join(r["groups"])),
    ("Highest severity", 12, lambda r: r.get("severity", "")),
    ("Findings", 48, lambda r: _findings(r.get("findings") or {})),
]
POLICY_HEADERS = [c[0] for c in POLICY_COLUMNS]


def finding_columns(name: str = "Role") -> list[Column]:
    """Columns of the Findings sheet; ``name`` heads the first one (Role, User, Policy)."""
    return [
        (name, 32, lambda f: f["role"]),
        ("Severity", 10, lambda f: f["severity"]),
        ("Finding", 34, lambda f: f["title"]),
        ("Rule", 26, lambda f: f["rule"]),
        ("Where", 28, lambda f: f["where"]),
        ("Statements", 18, lambda f: ", ".join(f["statements"])),
        ("Why", 60, lambda f: f["why"]),
        ("Suggested fix", 60, lambda f: f["fix"]),
    ]


FINDING_COLUMNS = finding_columns()
FINDING_HEADERS = [c[0] for c in FINDING_COLUMNS]


def _put(ws: Any, row: int, col: int, value: Any) -> None:
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # never a formula
    elif isinstance(value, datetime):
        cell.number_format = DATE_FORMAT


def _table(ws: Any, columns: list[Column], rows: Sequence[Mapping[str, Any]]) -> None:
    """A bold, frozen header row, one row per item and an autofilter over every column."""
    bold = Font(bold=True)
    for col, (title, width, _get) in enumerate(columns, start=1):
        _put(ws, 1, col, title)
        ws.cell(row=1, column=col).font = bold
        ws.cell(row=1, column=col).fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, row in enumerate(rows, start=2):
        for col, (_title, _width, get) in enumerate(columns, start=1):
            _put(ws, r, col, get(row))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def roles_workbook(
    rows: Sequence[Mapping[str, Any]],
    filters: RoleFilters,
    *,
    account: str,
    taken_at: str,
    total: int,
    findings: Sequence[Mapping[str, Any]] | None = None,
    now: datetime | None = None,
) -> bytes:
    """The ``.xlsx`` file of ``rows`` (already filtered by ``filters``) and the
    ``findings`` (dicts of :func:`rolelens.insights.findings_by_role`) of those rows."""
    return _workbook(
        "Roles", COLUMNS, "Role", rows, filters, account, taken_at, total, findings, now
    )


def users_workbook(
    rows: Sequence[Mapping[str, Any]],
    filters: UserFilters,
    *,
    account: str,
    taken_at: str,
    total: int,
    findings: Sequence[Mapping[str, Any]] | None = None,
    now: datetime | None = None,
) -> bytes:
    """The Users page as ``.xlsx``: rows of :func:`rolelens.queries.user_list`."""
    return _workbook(
        "Users", USER_COLUMNS, "User", rows, filters, account, taken_at, total, findings, now
    )


def policies_workbook(
    rows: Sequence[Mapping[str, Any]],
    filters: PolicyFilters,
    *,
    account: str,
    taken_at: str,
    total: int,
    findings: Sequence[Mapping[str, Any]] | None = None,
    now: datetime | None = None,
) -> bytes:
    """The Policies page as ``.xlsx``: rows of :func:`rolelens.queries.policy_list`."""
    return _workbook(
        "Policies", POLICY_COLUMNS, "Policy", rows, filters, account, taken_at, total, findings, now
    )


def _workbook(
    title: str,
    columns: list[Column],
    finding_name: str,
    rows: Sequence[Mapping[str, Any]],
    filters: RoleFilters | UserFilters | PolicyFilters,
    account: str,
    taken_at: str,
    total: int,
    findings: Sequence[Mapping[str, Any]] | None,
    now: datetime | None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    _table(ws, columns, rows)
    _table(wb.create_sheet("Findings"), finding_columns(finding_name), findings or [])
    bold = Font(bold=True)

    info = wb.create_sheet("Export")
    now = now or datetime.now(UTC)
    lines: list[tuple[str, Any]] = [
        ("Account", account),
        ("Collected (UTC)", _utc_naive(taken_at)),
        ("Exported (UTC)", _utc_naive(now.isoformat())),
        (title, f"{len(rows)} of {total}"),
        ("Findings", len(findings or [])),
        ("Filters", "none" if not filters.active else ""),
        *filters.describe(),
    ]
    for r, (label, value) in enumerate(lines, start=1):
        _put(info, r, 1, label)
        info.cell(row=r, column=1).font = bold
        _put(info, r, 2, value)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 48

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
